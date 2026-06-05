# -*- coding: utf-8 -*-
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
COVER = next(ROOT.glob("*封面.docx"))
OUT = ROOT / "果冻大逃亡_计算机图形学课程论文_无乱码.docx"


def set_font(run, name="宋体", size=11, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_paragraph(doc, text, indent=True, size=11, bold=False, color=None, align=None):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    p.paragraph_format.line_spacing = 1.45
    p.paragraph_format.space_after = Pt(5)
    if indent:
        p.paragraph_format.first_line_indent = Pt(22)
    r = p.add_run(text)
    set_font(r, "宋体", size, bold, color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_font(
        r,
        "黑体" if level <= 2 else "宋体",
        18 if level == 1 else 14 if level == 2 else 12,
        True,
        (31, 78, 121) if level == 1 else (50, 50, 50),
    )
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.65)
        p.paragraph_format.first_line_indent = Cm(-0.25)
        p.paragraph_format.line_spacing = 1.35
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run("• " + item)
        set_font(r, "宋体", 10.5)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell(cell, text, bold=False, fill=None, size=9.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run(str(text))
    set_font(r, "宋体", size, bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if fill:
        shade_cell(cell, fill)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    try:
        table.style = "Table Grid"
    except Exception:
        pass
    for i, header in enumerate(headers):
        set_cell(table.rows[0].cells[i], header, True, "D9EAF7", 10)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell(cells[i], value)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph()


def add_image(doc, image_name, caption, width=15.5):
    path = ROOT / image_name
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Cm(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    set_font(r, "宋体", 10, False, (90, 90, 90))


def setup_doc(doc):
    for section in doc.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.3)
        section.right_margin = Cm(2.3)
        footer = section.footer.paragraphs[0]
        footer.text = "果冻大逃亡双人迷宫动作游戏课程论文"
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in footer.runs:
            set_font(run, "宋体", 9)
    doc.styles["Normal"].font.name = "宋体"
    doc.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    doc.styles["Normal"].font.size = Pt(11)


def add_intro(doc):
    doc.add_page_break()
    add_heading(doc, "摘要", 1)
    for text in [
        "《果冻大逃亡》是一款基于 HTML5 Canvas 2D API、CSS 与 Vanilla JavaScript 实现的迷宫动作游戏。项目以二维网格地图为基础，围绕玩家移动、同色水果收集、敌人追踪、攻击判定、道具效果、关卡解锁、PK 竞速和合作通关等机制展开，尝试在网页环境中完成一套完整的小型图形学游戏系统。",
        "本文重点说明该游戏的研究目的与意义、总体设计思路、主要功能模块、类与对象关系、图形绘制方法、动画与交互实现、游戏规则、难度递进、运行效果以及与类似小游戏相比的改进点。通过该项目，可以较完整地体现计算机图形学课程中二维图形绘制、坐标变换、碰撞检测、动画循环、交互输入和界面组织等知识的实际应用。",
        "关键词：HTML5；Canvas 2D；JavaScript；迷宫游戏；AABB 碰撞检测；游戏动画；计算机图形学",
    ]:
        add_paragraph(doc, text)

    add_heading(doc, "目录", 1)
    for item in [
        "1 研究目的与意义",
        "2 开发环境与使用软件",
        "3 游戏总体设计",
        "4 场景与图形风格设计",
        "5 主要功能模块与类结构",
        "6 核心算法与关键实现",
        "7 角色、敌人与动画特效设计",
        "8 游戏规则、难度递进与关卡存档",
        "9 运行截图与界面说明",
        "10 与类似游戏的比较及创新点",
        "11 总结与展望",
        "附录：核心类、测试记录与改进计划",
    ]:
        add_paragraph(doc, item, indent=False)


SECTIONS = [
    (
        "1 研究目的与意义",
        [
            "本项目的主要目的是使用浏览器原生技术开发一款完整可运行的二维迷宫动作游戏。与单纯绘制静态图形不同，游戏项目需要同时处理图形表现、动画刷新、用户输入、物体碰撞、游戏状态管理和界面交互，因此能够更综合地体现计算机图形学与前端程序设计的结合。",
            "通过实现《果冻大逃亡》，可以训练将抽象的图形学知识转化为可交互程序的能力。例如，地图中的墙体和地面需要按照网格坐标映射到屏幕坐标；人物和敌人需要通过几何图形组合绘制成具有辨识度的小人；攻击、拾取和敌人技能需要用渐变、透明度、阴影和形状变化形成动态视觉反馈。",
            "从课程学习角度看，该游戏覆盖了计算机图形学中二维变换、颜色与渐变、层次绘制、透明度、阴影、动画帧刷新、空间碰撞等内容。学生不仅需要知道如何画一个图形，还要知道图形在时间轴上如何变化、如何响应玩家操作以及如何与其他对象发生关系。",
            "从软件设计角度看，游戏包含多种实体和状态，例如玩家、敌人、水果、道具、地图、关卡、暂停、结算等。如果所有逻辑都写在一个函数中，程序很快会变得混乱。因此本项目采用 ES6 Classes 组织实体对象，体现面向对象思想在小游戏开发中的应用。",
        ],
    ),
    (
        "2 开发环境与使用软件",
        [
            "本游戏采用纯网页技术开发，所有核心逻辑均运行在现代浏览器中，不需要安装额外游戏运行时。项目文件主要由 index.html、styles.css 和 game.js 三部分构成，其中 HTML 负责页面结构，CSS 负责界面布局与视觉样式，JavaScript 负责游戏逻辑和 Canvas 绘制。",
            "选择 Canvas 2D API 的原因是其适合进行二维图形绘制和动画刷新。游戏场景中的地面、墙体、人物、敌人、水果、道具和攻击特效都可以通过 Canvas 基本图元绘制完成，能够直接体现课程中的图形绘制知识。",
            "开发过程中使用浏览器进行运行测试，通过刷新页面观察主菜单、关卡选择、角色选择、游戏运行和结算界面的效果。由于游戏是纯前端项目，因此调试过程主要关注 JavaScript 逻辑是否报错、Canvas 是否正确绘制、键盘事件是否响应、碰撞检测是否准确以及 UI 状态是否互相覆盖。",
        ],
    ),
    (
        "3 游戏总体设计",
        [
            "《果冻大逃亡》定位为一款轻量级迷宫动作游戏。游戏的核心体验来自“收集、躲避、进门、竞争/合作”四个动作。玩家需要在由墙体和通道组成的地图中移动，收集与自身颜色对应的水果，躲避敌人追击，并在满足条件后进入出口大门。",
            "游戏支持三种模式：单人闯关、合作模式和 PK 模式。单人闯关强调个人操作和关卡进度；合作模式强调两名玩家分别收集水果并共同到达出口；PK 模式强调竞速和互相干扰。",
            "主菜单用于选择模式，关卡选择页用于读取存档并选择已解锁关卡，角色选择页用于选择角色和皮肤，游戏运行阶段负责实时交互，结算界面用于进入下一关、重试或返回菜单。游戏内按 Esc 可打开暂停菜单，暂停菜单中包含继续游戏、重开本关、查看规则和返回主菜单等功能。",
            "Game 类内部使用 state 字段管理状态，包括 menu、playing、paused、result 等。主循环在 state 为 playing 时更新对象位置和计时器，在 paused 或 result 状态下暂停逻辑更新。这样可以避免暂停时敌人继续移动或倒计时继续减少。",
        ],
    ),
    (
        "4 场景与图形风格设计",
        [
            "地图采用二维数组进行硬编码定义，数组中的 0 表示通道，1 表示墙体，2 表示出口大门。MazeMap 类负责读取矩阵、查找大门位置、判断瓦片类型以及提供随机通道位置。绘制时，程序将网格坐标乘以 tileSize，再加上居中偏移量，映射到 Canvas 屏幕坐标。",
            "场景视觉采用卡通沙盘风格。墙体不是简单的矩形，而是分为侧面和顶部两层，并通过圆角、渐变、阴影和高光表现立体感。地面则使用浅色渐变和少量装饰点，使画面不显得空。不同关卡会在颜色主题上发生变化，例如草地、冰原和沙地风格。",
            "人物没有使用外部精灵图，而是用 Canvas 基本图元组合绘制。drawLittlePerson 函数通过圆形头部、圆角矩形身体、线条手脚、帽子、头盔、兜帽等元素构成小人。不同角色使用不同配色和头饰，例如闪电队长、星际飞行员、糖果法师、青柠斥候。",
            "水果、道具和敌人也采用 Canvas 绘制。水果使用径向渐变、外圈描边和小叶子表现；道具使用圆形图标与文字标识；敌人则根据类型绘制不同外观，例如普通追兵、半透明幽灵、滑板怪、棱镜炮手和飞跃怪。",
        ],
    ),
    (
        "5 主要功能模块与类结构",
        [
            "为了降低复杂度，游戏核心采用面向对象方式组织。每一种主要实体都有相应的类或数据表，Game 类负责整体协调。Game 保存当前模式、关卡、地图、玩家数组、敌人数组、水果数组、道具数组、机关数组、时间和统计数据。",
            "MazeMap 类保存二维地图矩阵，Player 类处理玩家输入和状态，Enemy 类处理追踪和技能，Fruit 与 PowerUp 负责收集物和道具，绘制函数负责图形表现，UI 函数负责菜单和弹窗切换。",
            "Game.start 是每一局游戏的入口。它会根据当前模式和关卡创建 MazeMap，计算难度参数，生成玩家、敌人、水果、道具和机关，然后将状态切换为 playing。这个过程体现了游戏初始化的典型流程。",
            "Game.update 是逻辑更新函数，它只在 playing 状态下执行。该函数负责减少倒计时、更新玩家、更新敌人、刷新拾取特效、处理碰撞、处理合作救援和更新 HUD。Game.draw 则专门负责绘制，避免更新逻辑与图形表现混在一起。",
        ],
    ),
    (
        "6 核心算法与关键实现",
        [
            "游戏使用 requestAnimationFrame 构建主循环。浏览器会根据显示器刷新率调用 loop 函数，程序计算当前帧与上一帧的时间差 dt，并将 dt 传入 game.update(dt)。这样可以让对象移动与真实时间相关，而不是依赖固定帧数。",
            "项目中的墙体、玩家、敌人和攻击框主要使用 AABB 进行碰撞检测。AABB 的特点是计算简单，只需比较两个矩形在 x 和 y 方向上是否重叠。玩家移动时，会先根据目标位置构造一个矩形，再交给 MazeMap.collides 判断该矩形是否与墙体或未开启的大门重叠。",
            "敌人 AI 采用简化的网格寻路思想。敌人会周期性寻找最近玩家作为目标，然后在当前格子的上下左右四个方向中选择更接近目标的位置。虽然不是完整 A* 算法，但在小型迷宫中能提供足够自然的追踪效果。",
            "为了解决敌人在角落震荡的问题，Enemy 类记录 lastCell、lastPosition 和 stuckTimer。如果敌人长时间几乎没有移动，系统会强制重新选择方向。这样敌人不容易卡在墙角或来回抖动。",
        ],
    ),
    (
        "7 角色、敌人与动画特效设计",
        [
            "角色设计不仅区分颜色，还加入了性格描述和专属攻击特效。闪电队长热血、攻击为闪电电弧；星际飞行员冷静、攻击为长距离激光；糖果法师俏皮、攻击为星爆魔法；青柠斥候机灵、攻击为手雷爆炸。",
            "游戏动画主要通过时间函数和每帧重绘实现。例如人物移动时使用正弦函数产生轻微上下摆动，体现行走感；水果和道具使用脉冲缩放表现可拾取状态；攻击特效使用 attackTimer 计算进度，从而改变透明度、半径、线条宽度和粒子位置。",
            "Canvas 的 shadowBlur、globalAlpha、createLinearGradient 和 createRadialGradient 被大量用于特效表现。例如激光使用长条渐变和高亮描边，手雷使用星形爆炸轮廓和径向渐变，闪电使用折线和粒子表现电弧，星爆使用多角星图形表现魔法感。",
            "特殊敌人同样通过状态计时器实现技能动画。棱镜炮手先进入 charge 状态，画面显示较细的预警线，随后进入 fire 状态，激光变粗并具有伤害判定。飞跃怪进入 airborne 状态后会改变绘制高度，并在地面绘制阴影，表现短暂离地。",
        ],
    ),
    (
        "8 游戏规则、难度递进与关卡存档",
        [
            "单人闯关要求玩家在倒计时结束前吃完自己的水果并进入大门，被敌人抓到则失败。合作模式要求两名玩家各自吃完同色水果，大门开启后两人都进入大门才算通关。PK 模式要求两名玩家竞速，谁先吃完自己的水果并进入大门，谁获得胜利。",
            "PK 模式中被敌人抓到不会立刻失败，而是回到出生点并短暂眩晕。这样设计可以避免比赛因为一次碰撞立刻结束，同时又能产生明显惩罚。游戏难度随关卡递增，主要体现在水果数量、敌人数量、敌人速度、追踪范围、敌人技能类型和场景机关上。",
            "游戏使用 localStorage 保存进度，包括 bestSingleLevel、bestCoopLevel、pkWins 和 unlockedSkins。通关时 applyUnlocks 方法会更新最高关卡和皮肤解锁状态。进入关卡选择页时，buildLevelChoices 根据存档生成可点击关卡和锁定关卡。",
            "关卡选择规则是“最高通关关卡 + 1”。例如单人最高通关第 1 关，则第 1 关和第 2 关均可选择，第 3 关之后锁定。这样既保留闯关成就感，也避免玩家每次都必须从第一关重新开始。",
        ],
    ),
    (
        "10 与类似游戏的比较及创新点",
        [
            "传统迷宫游戏通常只要求玩家从入口走到出口，玩法重点是路径选择。本项目在迷宫基础上加入了水果收集、敌人 AI、攻击系统、道具系统、多人合作和 PK 竞速，使游戏目标更加多层次。玩家不仅要找路，还要考虑收集顺序、敌人位置、攻击时机和队友配合。",
            "从体验上看，本游戏受到合作闯关和轻派对游戏的启发，但它没有使用复杂的 3D 引擎，而是在 2D Canvas 中实现清晰的规则和可玩的反馈。与一些只做单人移动演示的课程项目相比，它包含模式切换、关卡存档、暂停菜单、皮肤解锁和多类型敌人，完整度更高。",
            "本项目的创新点包括：同色水果分工机制、合作与 PK 双目标、角色专属攻击、特殊敌人系统、关卡存档与皮肤解锁、Esc 暂停菜单等。这些设计使游戏从简单的图形演示扩展为具备完整流程的课程作品。",
        ],
    ),
    (
        "11 总结与展望",
        [
            "《果冻大逃亡》完成了一个基于 HTML5 Canvas 的二维迷宫动作游戏原型。项目实现了地图绘制、角色移动、AABB 碰撞、敌人 AI、攻击系统、水果拾取、道具效果、关卡选择、存档解锁、规则弹窗、暂停菜单、单人/合作/PK 多模式等功能。整体上，它不只是一个静态图形作业，而是一个具有完整流程和交互反馈的小游戏。",
            "通过该项目可以看出，Canvas 2D API 虽然接口较基础，但结合良好的对象设计和状态管理，仍然可以实现较丰富的游戏效果。项目中的人物、敌人、道具、墙体、门和攻击特效均由代码绘制，有助于理解图形绘制与动画变化的底层过程。",
            "未来可以继续扩展的方向包括：加入音效和背景音乐、增加更多地图、使用更完善的寻路算法、加入联网功能、设计更多道具和机关、增加移动端触控支持、优化敌人平衡、增加排行榜和成就系统等。如果继续迭代，本项目可以发展成一款更完整的网页休闲合作游戏。",
        ],
    ),
]


def add_main_sections(doc):
    doc.add_page_break()
    page_break_after = {"2 开发环境与使用软件", "4 场景与图形风格设计", "6 核心算法与关键实现", "8 游戏规则、难度递进与关卡存档"}
    for title, paragraphs in SECTIONS:
        add_heading(doc, title, 1)
        add_heading(doc, title.split(" ", 1)[0] + ".1 章节说明", 2)
        for text in paragraphs:
            add_paragraph(doc, text)
        if title in page_break_after:
            doc.add_page_break()


def add_tables(doc):
    add_heading(doc, "2.3 开发环境表", 2)
    add_table(
        doc,
        ["类别", "工具/技术", "作用"],
        [
            ["开发语言", "HTML5、CSS3、Vanilla JavaScript", "完成页面结构、视觉样式和游戏逻辑"],
            ["图形接口", "Canvas 2D API", "绘制地图、角色、敌人、道具、攻击特效和动画帧"],
            ["运行环境", "Microsoft Edge / Google Chrome", "直接打开 index.html 即可运行游戏"],
            ["代码组织", "ES6 Classes", "封装 Game、Player、Enemy、Fruit、PowerUp、MazeMap 等实体"],
        ],
        [2.8, 4.8, 8.2],
    )
    add_heading(doc, "5.3 主要模块职责表", 2)
    add_table(
        doc,
        ["模块", "主要类/结构", "职责说明"],
        [
            ["游戏控制模块", "Game", "管理模式、关卡、计时、主循环、胜负判断、HUD、暂停和结算"],
            ["地图模块", "MazeMap", "保存二维地图矩阵，判断墙体、通道、大门，生成随机通道位置"],
            ["玩家模块", "Player", "处理移动、朝向、攻击、护盾、磁吸、加速、倒地与复活"],
            ["敌人模块", "Enemy + ENEMY_TYPES", "实现不同敌人类型、追踪 AI、穿墙、滑板冲刺、激光和飞跃"],
            ["收集物模块", "Fruit", "根据玩家颜色生成水果，检测玩家拾取"],
            ["道具模块", "PowerUp", "实现冰冻、护盾、磁吸、加速等效果"],
        ],
        [2.8, 3.8, 8.8],
    )
    add_heading(doc, "7.3 敌人类型表", 2)
    add_table(
        doc,
        ["敌人类型", "外观特点", "能力特点", "出现阶段"],
        [
            ["果冻追兵", "棕色基础敌人", "普通追踪玩家", "第1关起"],
            ["穿墙幽灵", "半透明浅蓝色外观", "可越过墙体追踪", "第2关起"],
            ["滑板快客", "绿色身体与滑板", "短时间高速冲刺", "第3关起"],
            ["棱镜炮手", "粉色炮手外观", "蓄力后发射贯穿地图的激光", "第4关起"],
            ["飞跃怪", "紫色翅膀外观", "短暂飞起并落点突袭", "第5关起"],
        ],
        [2.6, 3.6, 6.2, 2.8],
    )


def add_screenshots(doc):
    doc.add_page_break()
    add_heading(doc, "9 运行截图与界面说明", 1)
    add_paragraph(doc, "以下截图均来自项目文件夹“游戏截图”中的实际运行画面，展示了游戏从主菜单、规则查看、关卡选择、角色选择到正式战斗、胜利结算和暂停菜单的完整流程。截图能够说明游戏不仅完成了 Canvas 场景绘制，也实现了多界面状态切换、关卡存档、角色解锁、道具反馈、特殊敌人与弹窗交互等功能。")
    screenshots = [
        ("游戏截图/屏幕截图 2026-06-04 160954.png", "图9-1 主菜单界面：标题、三种游戏模式、控制提示和“查看游戏规则”按钮集中显示，背景保留地图预览，避免纯空白菜单。"),
        ("游戏截图/屏幕截图 2026-06-04 161021.png", "图9-2 游戏规则界面：用弹窗方式说明单人、合作、PK、敌人、道具和角色特效规则，解决首页信息拥挤的问题。"),
        ("游戏截图/屏幕截图 2026-06-04 161036.png", "图9-3 单人闯关选关界面：显示已通关最高关、当前可选择关卡和未解锁关卡，体现 localStorage 存档机制。"),
        ("游戏截图/屏幕截图 2026-06-04 161100.png", "图9-4 单人角色选择界面：四名角色拥有不同性格、颜色皮肤和专属攻击招式，玩家确认后进入游戏。"),
        ("游戏截图/屏幕截图 2026-06-04 161117.png", "图9-5 单人游戏运行界面：玩家、同色水果、加速道具、护盾、敌人、出口大门和 HUD 同时呈现。"),
        ("游戏截图/屏幕截图 2026-06-04 161200.png", "图9-6 双人角色选择界面：两名玩家分别选择角色，适用于合作模式和 PK 模式。"),
        ("游戏截图/屏幕截图 2026-06-04 161233.png", "图9-7 PK 模式选关界面：PK 模式同样支持关卡选择和难度推进，胜利局数会影响解锁进度。"),
        ("游戏截图/屏幕截图 2026-06-04 161254.png", "图9-8 PK 模式运行界面：两名玩家同时收集水果，特殊敌人释放贯穿地图的激光，增强竞技干扰。"),
        ("游戏截图/屏幕截图 2026-06-04 162038.png", "图9-9 PK 胜利结算界面：玩家完成水果目标并进入大门后触发胜利弹窗，可进入下一关、重试或返回主菜单。"),
        ("游戏截图/屏幕截图 2026-06-04 162053.png", "图9-10 暂停菜单界面：按 Esc 可暂停游戏，提供继续游戏、重新本关、查看规则和返回主菜单等操作。"),
    ]
    for image_name, caption in screenshots:
        add_image(doc, image_name, caption, width=14.2)


def add_appendix(doc):
    appendices = [
        ("附录A：关键代码设计说明", [
            "本项目的核心代码集中在 game.js 中。虽然文件较长，但整体结构遵循“数据配置—工具函数—实体类—绘制函数—UI 事件”的顺序。开头部分定义了 TILE、MODE、POWERUPS、ENEMY_TYPES、CHARACTERS 和 SKINS 等常量，这些常量相当于游戏的配置层。",
            "Game.start 是每一局游戏的入口。它会根据当前模式和关卡创建 MazeMap，计算难度参数，生成玩家、敌人、水果、道具和机关，然后将状态切换为 playing。这个过程体现了游戏初始化的典型流程：先清空旧数据，再构建新场景，最后交给主循环执行。",
            "Game.update 是逻辑更新函数，它只在 playing 状态下执行。该函数负责减少倒计时、更新玩家、更新敌人、刷新拾取特效、处理碰撞、处理合作救援和更新 HUD。Game.draw 则专门负责绘制，避免更新逻辑与图形表现混在一起。",
        ]),
        ("附录B：算法流程说明", [
            "玩家移动流程可以概括为：读取键盘状态，得到水平和垂直方向输入；如果同时按下两个方向键，则进行归一化，避免斜向移动速度变快；根据移动方向更新角色朝向；计算速度，叠加加速道具、冰面、泥地等影响；分别尝试 x 方向和 y 方向移动；如果目标矩形与墙体或关闭的大门碰撞，则取消该方向移动。",
            "敌人追踪流程可以概括为：周期性寻找最近的有效玩家；如果目标距离超过丢失范围，则取消追踪；根据当前格子的上下左右四个候选方向评分；追踪状态下优先选择离目标更近的格子；巡逻状态下优先保持原方向并加入随机扰动；如果长时间没有移动，则强制换方向。",
            "攻击判定流程可以概括为：玩家按下攻击键后，如果冷却结束，则设置 attackTimer；在 attackTimer 大于 0 时，根据角色朝向和攻击类型生成攻击矩形；遍历敌人数组，如果攻击矩形与敌人碰撞矩形重叠，则击败敌人并生成爆炸或命中特效。",
        ]),
        ("附录C：测试记录与问题修正", [
            "在开发过程中，曾出现水果路过不被拾取的问题。修正方法是将水果拾取从矩形碰撞调整为距离检测，并扩大玩家与水果的接触阈值，同时确保水果只生成在通道上。这样玩家经过水果附近时能够稳定触发拾取反馈。",
            "敌人早期版本会在角落发生震荡。修正方法是在 Enemy 类中加入 lastPosition 和 stuckTimer，如果敌人一定时间内位移过小，就认为其被卡住，强制重新选择方向。该方法虽然简单，但对小型网格地图有明显改善。",
            "PK 模式中曾出现按空格导致水果和敌人位置突变的问题。原因是浏览器按钮焦点仍停留在页面按钮上，按空格会触发按钮点击。修正方法是在进入游戏时移除焦点，并在 keydown 中阻止 Space 默认行为。",
        ]),
        ("附录D：课堂演示步骤", [
            "课堂演示时可以按以下顺序进行：首先打开 index.html，展示主菜单和游戏规则弹窗；然后进入单人模式，展示关卡选择和角色选择；进入游戏后演示移动、拾取水果、攻击敌人和使用道具；接着按 Esc 展示暂停菜单；最后演示合作或 PK 模式，说明不同模式的胜利条件。",
            "演示重点不应只放在“能不能玩”，还应说明图形学实现方式。例如指出墙体、地面、人物、敌人和特效都是 Canvas 绘制；说明 requestAnimationFrame 如何驱动画面更新；说明 AABB 如何实现碰撞；说明 localStorage 如何保存关卡进度。",
            "如果时间允许，可以展示高关卡中的特殊敌人：穿墙幽灵、滑板快客、棱镜炮手和飞跃怪。它们能体现游戏机制扩展能力，也能说明项目不是简单的静态地图行走。",
        ]),
        ("附录E：课程知识点对应关系", [
            "二维坐标变换对应地图格子到屏幕像素的转换；颜色模型和渐变对应墙体、水果、道具和攻击特效的绘制；透明度和阴影对应光效、幽灵、拾取反馈和层次表现；动画循环对应 requestAnimationFrame；交互输入对应键盘事件；空间关系对应 AABB 和距离碰撞检测。",
            "该项目还体现了图形绘制与软件工程之间的关系。如果只有图形绘制而没有对象封装，游戏很难维护；如果只有逻辑而没有视觉反馈，玩家又难以理解当前状态。因此课程项目需要兼顾图形表现、交互反馈和代码结构。",
            "从学习效果看，本项目让抽象知识落到了具体功能中。例如“渐变”不再只是画一个渐变矩形，而是用于表现水果体积、墙体高光和激光能量；“碰撞检测”不再只是数学公式，而是决定玩家能否穿墙、能否吃到水果、能否被敌人抓到。",
        ]),
        ("附录F：后续优化计划", [
            "后续可以加入音效系统，为拾取、攻击、开门、失败和通关配置不同音效，增强反馈。也可以加入背景音乐，使不同主题关卡具有更明显的氛围。音频可以通过 HTMLAudioElement 或 Web Audio API 实现。",
            "地图方面可以设计可视化地图编辑器，将二维数组编辑过程图形化。这样后续制作关卡时不需要手动修改矩阵，也可以保存更多地图数据。敌人 AI 方面可以进一步使用 A* 寻路，使敌人在复杂地图中能更准确地追踪玩家。",
            "如果继续拓展为完整游戏，还可以加入移动端虚拟摇杆、排行榜、成就系统、联网双人、更多人物皮肤、更多道具和更多陷阱。通过这些扩展，项目可以从课程作业原型发展为较完整的网页休闲小游戏。",
        ]),
    ]
    for title, paragraphs in appendices:
        doc.add_page_break()
        add_heading(doc, title, 1)
        for text in paragraphs:
            add_paragraph(doc, text)


def build():
    doc = Document(str(COVER))
    setup_doc(doc)
    add_intro(doc)
    add_main_sections(doc)
    add_tables(doc)
    add_screenshots(doc)
    add_appendix(doc)
    doc.add_page_break()
    add_heading(doc, "参考资料", 1)
    for item in [
        "HTML Living Standard：Canvas 元素与浏览器事件模型相关说明。",
        "MDN Web Docs：CanvasRenderingContext2D、requestAnimationFrame、KeyboardEvent、localStorage 文档。",
        "计算机图形学课程资料：二维图形绘制、坐标变换、颜色模型、动画与交互基础。",
        "经典迷宫游戏与合作闯关游戏体验分析：地图通道设计、多人协作、障碍与奖励机制。",
    ]:
        add_paragraph(doc, item, indent=False)
    doc.save(str(OUT))
    print(str(OUT))


if __name__ == "__main__":
    build()
